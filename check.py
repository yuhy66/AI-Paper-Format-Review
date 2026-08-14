from docx import Document

class Checker1:
    def __init__(self, file_path="test.docx"):
        self.file_path = file_path
        self.doc = None
        self.result_list = []

    def load_doc(self):
        try:
            self.doc = Document(self.file_path)
            self.result_list.append("✅ Word文档读取成功")
            return True
        except Exception as e:
            self.result_list.append(f"❌ 文件读取失败：{str(e)}")
            return False

    def check_cover(self):
        """一、封面检查"""
        self.result_list.append("\n===== 【封面单元检查】 =====")
        cover_text = []
        for para in self.doc.paragraphs[:26]:
            t = para.text.strip()
            if t:
                cover_text.append(t)
        all_cov = "".join(cover_text)

        # 封面必填项
        must_cover = ["论文题目", "院系", "专业", "姓名", "学号", "指导教师", "提交日期"]
        for item in must_cover:
            if item in all_cov:
                self.result_list.append(f"✅ 封面包含字段：{item}")
            else:
                self.result_list.append(f"⚠️ 封面缺失：{item}，请核对模板")

        # 提取题目粗判长度（≤25字）
        title_text = ""
        for p in self.doc.paragraphs[:12]:
            pt = p.text.strip()
            if 5 <= len(pt) <= 35:
                title_text = pt
                break
        if title_text:
            t_len = len(title_text)
            self.result_list.append(f"检测封面题目文字长度：{t_len}字")
            if t_len <= 25:
                self.result_list.append("✅ 题目字数≤25，符合要求")
            else:
                self.result_list.append("⚠️ 封面题目超过25个字！")
        else:
            self.result_list.append("⚠️ 未能识别封面论文题目")
        self.result_list.append("💡备注：字体字号、居中位置无法自动检测，需人工对照模板检查")

    def check_title_page(self):
        """二、扉页单元检查"""
        self.result_list.append("\n===== 【扉页单元检查】 =====")
        page2_text = []
        for para in self.doc.paragraphs[26:56]:
            t = para.text.strip()
            if t:
                page2_text.append(t)
        fp_text = "".join(page2_text)

        fp_check = ["中文题目", "英文题目", "姓名", "学号", "院系", "专业", "指导教师", "提交日期"]
        for w in fp_check:
            if w in fp_text:
                self.result_list.append(f"✅ 扉页检测到：{w}")
            else:
                self.result_list.append(f"⚠️ 扉页缺少：{w}")
        self.result_list.append("💡备注：扉页中英文标题字体、字号、居中无法代码检测，人工核验")

    def check_abstract(self):
        """三、摘要单元（中文摘要+关键词 + 英文摘要）"""
        self.result_list.append("\n===== 【中文摘要 & 关键词检查】 =====")
        abs_content = ""
        keyword_line = ""
        state = 0
        for para in self.doc.paragraphs:
            p_txt = para.text.strip()
            if not p_txt:
                continue
            if "摘要" in p_txt and state == 0:
                state = 1
                continue
            if state == 1:
                if "关键词" in p_txt:
                    keyword_line = p_txt
                    state = 2
                    break
                abs_content += p_txt

        # 中文摘要字数 300‑500
        abs_len = len(abs_content)
        self.result_list.append(f"中文摘要字符统计：{abs_len}字")
        if 300 <= abs_len <= 500:
            self.result_list.append("✅ 摘要字数300‑500，符合规范")
        else:
            self.result_list.append("⚠️ 摘要字数不在300‑500区间")

        # 关键词检查
        if keyword_line:
            self.result_list.append(f"读取关键词行：{keyword_line}")
            kw_part = keyword_line.replace("关键词：", "").replace("关键词:", "")
            kw_list = [k.strip() for k in kw_part.split("，") if k.strip()]
            kw_num = len(kw_list)
            self.result_list.append(f"关键词数量：{kw_num}个")
            if 3 <= kw_num <= 5:
                self.result_list.append("✅ 关键词数量3‑5个，符合要求")
            else:
                self.result_list.append("⚠️ 关键词数量应为3‑5个")
            self.result_list.append("💡关键词排序（范围从大到小）、字体加粗需要人工检查")
        else:
            self.result_list.append("⚠️ 没有找到关键词段落")

        self.result_list.append("\n===== 【英文摘要检查提示】 =====")
        self.result_list.append("💡中英文摘要内容一致性、单独分页、英文字体格式需要人工核对")

    def save_report_txt(self):
        """导出检查报告保存到txt文件"""
        with open("check_result.txt", "w", encoding="utf‑8") as f:
            for line in self.result_list:
                f.write(line + "\n")
        self.result_list.append("\n📄检测报告已保存至 check_result.txt")

    def run_all_check(self):
        print("===== Checker1 封面‑扉页‑摘要检测启动 =====")
        if not self.load_doc():
            for res in self.result_list:
                print(res)
            return
        self.check_cover()
        self.check_title_page()
        self.check_abstract()
        self.save_report_txt()
        for line in self.result_list:
            print(line)


if __name__ == "__main__":
    checker = Checker1("test.docx")
    checker.run_all_check()